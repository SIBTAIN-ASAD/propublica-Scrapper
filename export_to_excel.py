import sqlite3
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


OUT_DIR = Path("propublica_output")
DB_PATH = OUT_DIR / "propublica_scrape.sqlite"
EXCEL_PATH = OUT_DIR / "propublica_all_states.xlsx"

STATES = [
    "CA", "NY", "TX", "FL", "PA", "IL", "OH", "MA", "VA", "NC",
    "MI", "NJ", "WA", "CO", "DC", "MN", "GA", "MD", "WI", "IN",
    "MO", "TN", "OR", "CT", "AZ", "SC", "IA", "LA", "AL", "KY",
    "OK", "DE", "KS", "UT", "NE", "NV", "NM", "AR", "MS", "ME",
    "WV", "HI", "MT", "VT", "NH", "RI", "ID", "AK", "SD", "ND",
    "WY", "PR", "VI", "GU", "MP", "PW"
]

HEADERS = [
    "State",
    "Page",
    "Org Name",
    "Address",
    "Employer ID",
    "Phone",
    "Email",
    "Org Website",
    "Org URL",
    "Filing URL",
    "IRS990 URL",
    "Status",
    "Error",
]

WARNING_PREFIX = "WARNING:"


def row_status(error):
    if not error:
        return "OK"
    if str(error).startswith(WARNING_PREFIX):
        return "WARNING"
    return "ERROR"


def autosize_columns(ws):
    max_width = 45

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        longest = 0

        for cell in col:
            value = str(cell.value) if cell.value else ""
            longest = max(longest, len(value))

        ws.column_dimensions[col_letter].width = min(longest + 3, max_width)


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    light_fill = PatternFill("solid", fgColor="D9EAF7")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    error_fill = PatternFill("solid", fgColor="F4CCCC")
    status_col = HEADERS.index("Status") + 1

    thin_border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row in ws.iter_rows(min_row=2):
        status = row[status_col - 1].value

        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            if status == "ERROR":
                cell.fill = error_fill
            elif status == "WARNING":
                cell.fill = warning_fill
            elif cell.row % 2 == 0:
                cell.fill = light_fill

    autosize_columns(ws)


def get_summary_stats(cur):
    cur.execute("SELECT COUNT(*) FROM results")
    total = cur.fetchone()[0]

    cur.execute(
        "SELECT COUNT(*) FROM results WHERE error IS NULL OR error = ''"
    )
    ok_count = cur.fetchone()[0]

    cur.execute(
        "SELECT COUNT(*) FROM results WHERE error LIKE ?",
        (f"{WARNING_PREFIX}%",),
    )
    warned_count = cur.fetchone()[0]

    errored_count = total - ok_count - warned_count

    cur.execute("""
        SELECT
            state,
            COUNT(*) AS total,
            SUM(CASE WHEN error IS NULL OR error = '' THEN 1 ELSE 0 END) AS ok_count,
            SUM(CASE WHEN error LIKE ? THEN 1 ELSE 0 END) AS warned_count,
            SUM(
                CASE
                    WHEN error IS NOT NULL
                     AND error != ''
                     AND error NOT LIKE ?
                    THEN 1 ELSE 0
                END
            ) AS errored_count
        FROM results
        GROUP BY state
        ORDER BY state
    """, (f"{WARNING_PREFIX}%", f"{WARNING_PREFIX}%"))

    per_state = {
        row[0]: {
            "total": row[1],
            "ok": row[2],
            "warned": row[3],
            "errored": row[4],
        }
        for row in cur.fetchall()
    }

    return {
        "total": total,
        "ok": ok_count,
        "warned": warned_count,
        "errored": errored_count,
        "per_state": per_state,
    }


def create_summary_sheet(wb, summary_stats):
    ws = wb.create_sheet("Summary", 0)

    ws.append(["Metric", "Count"])
    ws.append(["Total records", summary_stats["total"]])
    ws.append(["Successful (OK)", summary_stats["ok"]])
    ws.append(["Warnings", summary_stats["warned"]])
    ws.append(["Errors", summary_stats["errored"]])
    ws.append([])

    ws.append(["State", "Total", "OK", "Warnings", "Errors"])
    for state in STATES:
        counts = summary_stats["per_state"].get(state, {
            "total": 0,
            "ok": 0,
            "warned": 0,
            "errored": 0,
        })
        ws.append([
            state,
            counts["total"],
            counts["ok"],
            counts["warned"],
            counts["errored"],
        ])

    header_fill = PatternFill("solid", fgColor="7030A0")
    header_font = Font(color="FFFFFF", bold=True)
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    error_fill = PatternFill("solid", fgColor="F4CCCC")
    error_font = Font(color="9C0006")

    for cell in ws["A1:B1"][0]:
        cell.fill = header_fill
        cell.font = header_font

    state_header_row = 8
    for cell in ws[state_header_row]:
        cell.fill = header_fill
        cell.font = header_font

    for row in ws.iter_rows(min_row=state_header_row + 1):
        if not row[0].value:
            continue

        warnings = row[3].value or 0
        errors = row[4].value or 0

        for cell in row:
            cell.alignment = Alignment(horizontal="center")
            cell.border = Border(
                left=Side(style="thin", color="D9E2F3"),
                right=Side(style="thin", color="D9E2F3"),
                top=Side(style="thin", color="D9E2F3"),
                bottom=Side(style="thin", color="D9E2F3"),
            )

        if warnings:
            row[3].fill = warning_fill
        if errors:
            row[4].fill = error_fill
            row[4].font = error_font

    ws.freeze_panes = "A2"
    autosize_columns(ws)


def export_to_excel():
    if not DB_PATH.exists():
        raise FileNotFoundError(f"SQLite database not found: {DB_PATH}")

    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    summary_stats = get_summary_stats(cur)

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    for state in STATES:
        print(f"[EXPORT] Creating sheet for {state}")

        ws = wb.create_sheet(title=state)
        ws.append(HEADERS)

        cur.execute("""
            SELECT
                state,
                page,
                org_name,
                address,
                employer_id,
                phone,
                email,
                org_website,
                org_url,
                filing_url,
                irs990_url,
                error
            FROM results
            WHERE state = ?
            ORDER BY page, org_name
        """, (state,))

        rows = cur.fetchall()

        for row in rows:
            error = row[-1]
            ws.append([*row[:-1], row_status(error), error])

        if rows:
            table_ref = f"A1:{get_column_letter(len(HEADERS))}{len(rows) + 1}"
            table = Table(displayName=f"Table_{state}", ref=table_ref)

            style = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            table.tableStyleInfo = style
            ws.add_table(table)

        style_sheet(ws)

        counts = summary_stats["per_state"].get(state, {"total": 0, "ok": 0, "warned": 0, "errored": 0})
        print(
            f"[DONE] {state}: {counts['total']} rows "
            f"(ok={counts['ok']} warned={counts['warned']} errored={counts['errored']})"
        )

    create_summary_sheet(wb, summary_stats)

    wb.save(EXCEL_PATH)
    conn.close()

    print(f"\n[SAVED] Excel exported successfully:")
    print(EXCEL_PATH)
    print(
        f"[SUMMARY] total={summary_stats['total']} "
        f"ok={summary_stats['ok']} "
        f"warned={summary_stats['warned']} "
        f"errored={summary_stats['errored']}"
    )


if __name__ == "__main__":
    export_to_excel()
